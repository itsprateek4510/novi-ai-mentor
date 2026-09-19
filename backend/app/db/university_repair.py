"""Repair the QS-imported universities table so the explorer is logical.

Fixes applied (idempotent, safe to re-run):
  - Re-derive subject / course / ranking / courses from the QS xlsx using a
    known mapping from (truncated) workbook sheet names to canonical subjects.
  - `subject`         -> canonical primary subject slug (best-ranked subject).
  - `course`          -> clean label of the primary subject (never truncated).
  - `courses`         -> list of clean labels for every subject the uni is ranked in.
  - `rankings`        -> {subject_slug: best_rank} so subject-filtered order is correct.
  - `ranking`         -> best rank across all subjects (ties broken by priority).
  - `strengths`       -> top subject labels for QS-derived rows (nicer detail page).
  - Country aliases   -> normalized (USA->..., UK->..., etc.).
  - Curated duplicate -> QS rows that are the same institution as a curated row
    are merged into the curated row and the QS duplicate is deleted.

Usage:  cd backend && ../venv/bin/python -m app.db.university_repair
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from sqlalchemy import select, text  # noqa: E402

from app.core.database import SessionLocal, engine  # noqa: E402
from app.models.university import University  # noqa: E402

QS_XLSX = Path(
    "/Users/itsprateek4510/Downloads/"
    "QS World University Rankings by Subject 2026 - Public Results v1.4 (qs.com)_26.xlsx"
)

# Truncated workbook sheet name -> (canonical subject slug, clean display label).
# The QS public download truncates sheet/tab names to 31 chars.
SHEET_MAP = {
    "Archaeology": ("arts", "Archaeology"),
    "Architecture _ Built Environmen": ("architecture", "Architecture & Built Environment"),
    "Art & Design": ("arts", "Art & Design"),
    "Classics & Ancient History": ("arts", "Classics & Ancient History"),
    "English Language & Literature": ("arts", "English Language & Literature"),
    "History_Subject": ("arts", "History"),
    "History of Art": ("arts", "History of Art"),
    "Linguistics": ("arts", "Linguistics"),
    "Modern Languages": ("arts", "Modern Languages"),
    "Music": ("arts", "Music"),
    "Performing Arts": ("arts", "Performing Arts"),
    "Philosophy": ("arts", "Philosophy"),
    "Theology, Divinity & Religious ": ("arts", "Theology, Divinity & Religious Studies"),
    "Computer Science & Information ": ("computer-science", "Computer Science & Information Systems"),
    "Data Science and Artificial Int": ("data-science", "Data Science & Artificial Intelligence"),
    "Engineering - Chemical": ("engineering", "Engineering - Chemical"),
    "Engineering - Civil & Structura": ("engineering", "Engineering - Civil & Structural"),
    "Engineering - Electrical & Elec": ("engineering", "Engineering - Electrical & Electronic"),
    "Engineering - Mechanical, Aeron": ("engineering", "Engineering - Mechanical, Aeronautical & Manufacturing"),
    "Engineering - Mineral & Mining": ("engineering", "Engineering - Mineral & Mining"),
    "Petroleum Engineering": ("engineering", "Petroleum Engineering"),
    "Agriculture & Forestry": ("science", "Agriculture & Forestry"),
    "Anatomy & Physiology": ("science", "Anatomy & Physiology"),
    "Biological Sciences": ("science", "Biological Sciences"),
    "Dentistry": ("medicine", "Dentistry"),
    "Medicine": ("medicine", "Medicine"),
    "Nursing": ("medicine", "Nursing"),
    "Pharmacy & Pharmacology": ("medicine", "Pharmacy & Pharmacology"),
    "Psychology": ("science", "Psychology"),
    "Veterinary Science": ("medicine", "Veterinary Science"),
    "Chemistry": ("science", "Chemistry"),
    "Earth & Marine Sciences": ("science", "Earth & Marine Sciences"),
    "Environmental Sciences": ("science", "Environmental Sciences"),
    "Geography": ("science", "Geography"),
    "Geology": ("science", "Geology"),
    "Geophysics": ("science", "Geophysics"),
    "Materials Science": ("engineering", "Materials Science"),
    "Mathematics": ("science", "Mathematics"),
    "Physics & Astronomy": ("science", "Physics & Astronomy"),
    "Accounting & Finance": ("business", "Accounting & Finance"),
    "Anthropology": ("arts", "Anthropology"),
    "Business & Management Studies": ("business", "Business & Management Studies"),
    "Communication & Media Studies": ("arts", "Communication & Media Studies"),
    "Development Studies": ("economics", "Development Studies"),
    "Economics & Econometrics": ("economics", "Economics & Econometrics"),
    "Education": ("arts", "Education"),
    "Hospitality & Leisure Managemen": ("business", "Hospitality & Leisure Management"),
    "Law": ("law", "Law"),
    "Library & Information Managemen": ("science", "Library & Information Management"),
    "Marketing": ("business", "Marketing"),
    "Politics & International Studie": ("economics", "Politics & International Studies"),
    "Social Policy & Administration": ("economics", "Social Policy & Administration"),
    "Sociology": ("arts", "Sociology"),
    "Sports-related Subjects": ("science", "Sports-related Subjects"),
    "Statistics & Operational Resear": ("science", "Statistics & Operational Research"),
}
SKIP_SHEETS = {
    "Index", "methodology", "Arts & Humanities", "Engineering & Technology",
    "Life Sciences & Medicine", "Natural Sciences", "Social Sciences & Management",
}
SUBJECT_PRIORITY = [
    "computer-science", "data-science", "engineering", "architecture", "medicine",
    "law", "business", "economics", "science", "arts",
]
SLUG_LABEL = {meta[0]: meta[1] for meta in SHEET_MAP.values()}
LABEL_SLUG = {meta[1]: meta[0] for meta in SHEET_MAP.values()}
COUNTRY_ALIASES = {
    "USA": "United States of America", "UK": "United Kingdom",
    "China": "China (Mainland)", "Hong Kong SAR": "Hong Kong SAR, China",
    "Macao": "Macao SAR, China", "Macao SAR": "Macao SAR, China",
    "Russia": "Russia", "South Korea": "South Korea",
    "Vietnam": "Viet Nam", "Czech Republic": "Czechia",
}
# Curated row name -> QS row names for the SAME institution (to merge + delete dup).
CURATED_QS_ALIASES = {
    "MIT": ["Massachusetts Institute of Technology (MIT)"],
    "UC Berkeley": ["University of California, Berkeley (UCB)"],
    "National University of Singapore": ["National University of Singapore (NUS)"],
    "IIT Bombay": ["Indian Institute of Technology Bombay (IITB)"],
    "IIT Delhi": ["Indian Institute of Technology Delhi (IITD)"],
    "BITS Pilani": ["Birla Institute of Technology and Science, Pilani"],
    "Stanford University": ["Stanford University"],
    "Harvard University": ["Harvard University"],
    "University of Oxford": ["University of Oxford"],
    "Imperial College London": ["Imperial College London"],
    "ETH Zurich": ["ETH Zurich"],
    "University of Waterloo": ["University of Waterloo"],
    "University of Cambridge": ["University of Cambridge"],
    "Carnegie Mellon University": ["Carnegie Mellon University"],
}
# Curated subject strings -> canonical slug style (keeps filter lists consistent).
CURATED_SUBJECT_MAP = {
    "computer science": "computer-science",
    "computer-science": "computer-science",
    "engineering": "engineering",
    "general": "general",
    "business & management": "business",
    "business": "business",
    "pharmacy": "medicine",
    "dentistry": "medicine",
    "data science": "data-science",
    "data-science": "data-science",
    "finance": "business",
}


def norm_country(country: str) -> str:
    return COUNTRY_ALIASES.get(country.strip(), country.strip())


def parse_xlsx() -> dict[tuple[str, str], dict]:
    """Return {(name_lower, country_norm): {name, country, subjects:{slug:best_rank},
    ranked:{label:best_rank}, labels:set}}."""
    wb = openpyxl.load_workbook(QS_XLSX, read_only=True, data_only=True)
    records: dict[tuple[str, str], dict] = {}
    for sheet in wb.sheetnames:
        if sheet in SKIP_SHEETS or sheet.startswith("arts & humanities"):
            continue
        meta = SHEET_MAP.get(sheet)
        if meta is None:
            continue
        slug, label = meta
        ws = wb[sheet]
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx < 4 or not row or len(row) < 4:
                continue
            rank, _prev, name, country = row[0], row[1], row[2], row[3]
            if name is None or not str(name).strip():
                continue
            name = str(name).strip()
            country = norm_country(str(country).strip() if country else "")
            key = (name.lower(), country)
            rec = records.get(key) or {
                "name": name, "country": country, "subjects": {}, "ranked": {}, "labels": set(), "slugs": set(),
            }
            rec["labels"].add(label)
            rec["slugs"].add(slug)
            if isinstance(rank, (int, float)):
                r = int(rank)
                label_key = (label, slug)
                if slug in SUBJECT_PRIORITY and ("subjects" not in rec or slug not in rec["subjects"] or r < rec["subjects"][slug]):
                    rec["subjects"][slug] = r
                if label not in rec["ranked"] or r < rec["ranked"][label]:
                    rec["ranked"][label] = r
            records[key] = rec
    wb.close()
    return records


def primary_subject(subjects: dict[str, int]) -> str:
    if not subjects:
        return "general"
    best = min(subjects.items(), key=lambda kv: (kv[1], SUBJECT_PRIORITY.index(kv[0]) if kv[0] in SUBJECT_PRIORITY else 99))
    return best[0]


def ensure_rankings_column() -> None:
    with engine.connect() as conn:
        cols = {r[0] for r in conn.execute(text("show columns from universities"))}
        if "rankings" not in cols:
            conn.execute(text("ALTER TABLE universities ADD COLUMN rankings JSON NULL"))
            conn.commit()
            print("[repair] added `rankings` column")


def main() -> None:
    ensure_rankings_column()
    qs = parse_xlsx()
    print(f"[repair] QS records parsed: {len(qs)}")

    db = SessionLocal()
    try:
        rows = list(db.scalars(select(University).order_by(University.id)))
        curated_by_name = {}
        qs_rows = {}
        for u in rows:
            if u.website:
                curated_by_name[u.name.strip().lower()] = u
            else:
                qs_rows[((u.name or "").strip().lower(), (u.country or "").strip())] = u

        updated = strengths_set = 0

        # 1) Merge QS data into curated rows + collect QS duplicates to delete.
        dup_ids: list[int] = []
        for curated_name, qs_names in CURATED_QS_ALIASES.items():
            curated = curated_by_name.get(curated_name.lower())
            if not curated:
                continue
            for qs_name in qs_names:
                dups = [
                    u for u in rows
                    if not u.website and u.name.strip().lower() == qs_name.lower()
                ]
                best_qs = qs.get((qs_name.lower(), norm_country(curated.country or "").strip()))
                if not dups and best_qs is None:
                    continue
                if best_qs and curated.country and norm_country(curated.country) == norm_country(best_qs["country"]):
                    if curated.subject in ("general", "") or CURATED_SUBJECT_MAP.get(curated.subject, curated.subject) not in best_qs["subjects"]:
                        primary = primary_subject(best_qs["subjects"])
                        ordered = sorted(best_qs["subjects"].items(), key=lambda kv: (kv[1], SUBJECT_PRIORITY.index(kv[0]) if kv[0] in SUBJECT_PRIORITY else 99))
                        for lead, lead_rank in ordered[:3]:
                            label = SLUG_LABEL.get(lead, lead)
                            if not any(lead in (c or "") for c in (curated.strengths or [])):
                                curated.strengths = (curated.strengths or []) + [f"Ranked #{lead_rank} in {label}"]
                        curated.subject = primary if curated.subject in ("general", "") or CURATED_SUBJECT_MAP.get(curated.subject, curated.subject) not in best_qs["subjects"] else CURATED_SUBJECT_MAP.get(curated.subject, curated.subject)
                curated.subject = CURATED_SUBJECT_MAP.get(curated.subject, curated.subject)
                if best_qs and best_qs["subjects"] and not curated.rankings:
                    curated.rankings = dict(best_qs["subjects"])
                    curated.courses = sorted({SLUG_LABEL.get(s, s) for s in curated.rankings})
                    if curated.ranking is None or curated.ranking > min(curated.rankings.values()):
                        curated.ranking = min(curated.rankings.values())
                    if curated.subject in ("general", "") or curated.subject not in curated.rankings:
                        curated.subject = primary_subject(curated.rankings)
                    if not curated.course or curated.course.strip() == "":
                        curated.course = SLUG_LABEL.get(curated.subject, curated.subject)
                    curated.country = norm_country(curated.country)
                for dup in dups:
                    dup_ids.append(dup.id)
                    merged = qs.get((dup.name.strip().lower(), (dup.country or "").strip()))
                    if merged and merged["subjects"]:
                        subjects = dict(merged["subjects"])
                        if curated.rankings:
                            subjects.update({k: v for k, v in curated.rankings.items() if k not in subjects})
                        primary = primary_subject(subjects)
                        curated.rankings = subjects
                        curated.courses = sorted({SLUG_LABEL.get(s, s) for s in subjects})
                        if curated.ranking is None or curated.ranking > min(subjects.values()):
                            curated.ranking = min(subjects.values())
                        if curated.subject in ("general", "") or curated.subject not in subjects:
                            curated.subject = primary
                        if not curated.course or curated.course.strip() == "":
                            curated.course = SLUG_LABEL.get(primary, primary)
                        curated.country = norm_country(curated.country)
                    updated += 1

        # 2) Repair all QS-derived rows.
        for u in rows:
            if u.id in dup_ids:
                continue
            if u.website:
                u.country = norm_country(u.country)
                u.subject = CURATED_SUBJECT_MAP.get(u.subject, u.subject)
                if u.rankings:
                    u.ranking = min(u.rankings.values())
                    u.courses = sorted({SLUG_LABEL.get(s, s) for s in u.rankings})
                continue
            u.country = norm_country(u.country)
            key = ((u.name or "").strip().lower(), u.country)
            rec = qs.get(key)
            if not rec:
                key = ((u.name or "").strip().lower(), (u.country or "").strip())
                rec = qs.get(key)
            if not rec:
                continue
            subjects = rec["subjects"]
            labels = sorted(rec["labels"])
            ranked = rec.get("ranked") or {}
            if not labels and subjects:
                labels = sorted({SLUG_LABEL.get(s, s) for s in subjects})
            u.courses = labels
            u.rankings = {s: r for s, r in subjects.items()}
            u.country = norm_country(u.country)
            if ranked:
                # course = the specific subject the university ranks best in.
                best_label = min(ranked.items(), key=lambda kv: (kv[1], SUBJECT_PRIORITY.index(LABEL_SLUG.get(kv[0], "arts")) if LABEL_SLUG.get(kv[0], "arts") in SUBJECT_PRIORITY else 99))[0]
                u.subject = LABEL_SLUG.get(best_label, u.subject or "general")
                u.course = best_label
                u.ranking = min(ranked.values())
            elif subjects:
                primary = primary_subject(subjects)
                u.subject = primary
                u.course = SLUG_LABEL.get(primary, primary)
                u.ranking = min(subjects.values())
            else:
                # No numeric rank in the public results for this uni. Derive the
                # primary subject from the sheets it is present in (priority order).
                slugs = sorted(
                    rec["slugs"],
                    key=lambda s: SUBJECT_PRIORITY.index(s) if s in SUBJECT_PRIORITY else 99,
                )
                primary = slugs[0] if slugs else CURATED_SUBJECT_MAP.get(u.subject, u.subject)
                u.subject = CURATED_SUBJECT_MAP.get(primary, primary)
                u.course = SLUG_LABEL.get(u.subject, labels[0] if labels else u.subject)
            ordered = sorted(ranked.items(), key=lambda kv: (kv[1], SUBJECT_PRIORITY.index(LABEL_SLUG.get(kv[0], "arts")) if LABEL_SLUG.get(kv[0], "arts") in SUBJECT_PRIORITY else 99))
            new_strengths = []
            for lead_label, lead_rank in ordered[:3]:
                new_strengths.append(f"Ranked #{lead_rank} in {lead_label}")
            u.strengths = new_strengths or (u.strengths or [])
            strengths_set += 1
            updated += 1

        if dup_ids:
            db.query(University).filter(University.id.in_(dup_ids)).delete(synchronize_session=False)
            print(f"[repair] deleted duplicate QS rows: {len(dup_ids)} (ids {sorted(dup_ids)})")

        db.commit()
        print(f"[repair] updated={updated} strengths_set={strengths_set} dup_deleted={len(dup_ids)}")
    finally:
        db.close()


if __name__ == "__main__":
    main()